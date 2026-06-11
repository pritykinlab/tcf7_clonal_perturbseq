#!/usr/bin/env python
"""Generate Figure 7 supplementary source-data tables."""

from __future__ import annotations

import argparse
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import load_workbook
import pandas as pd


# Required font settings for all manuscript figures
plt.rcParams["svg.fonttype"] = "none"
plt.rcParams["font.family"] = "sans-serif"
plt.rcParams["font.sans-serif"] = ["Arial"]
plt.rcParams["mathtext.default"] = "regular"


FIGURE7_FILES = {
    "7_1": "Fig7A_TCF7_validation_clean.csv",
    "7_2": "Fig7B_competitive_recovery_excluding_pd1_overlap_clean.csv",
    "7_3": "Fig7C_PD1_enrichment_plus_exp74_mannwhitney_clean.csv",
    "7_4": (
        "Fig7D_exp71_growth_no_pd1_forwardfilled_clean.csv",
        "Fig7D_exp71_survival_no_pd1_forwardfilled_clean.csv",
        "Fig7D_exp67_growth_pd1_clean.csv",
        "Fig7D_exp67_survival_pd1_clean.csv",
        "Fig7D_exp71_excluded_trim28_75pct.csv",
    ),
    "7_5": (
        "Fig7E_skin_distribution_3bin_clean.csv",
        "Fig7E_peritumor_vs_distant_pairs.csv",
        "Fig7E_3bin_omitted_adjacent_flanking_skin.csv",
    ),
}

SKIN_INPUT_TRIM28_FRACTIONS = {
    "Exp66": 0.523,
    "Exp74": 0.5352009921937696,
}

SKIN_EXCLUDED_EXPERIMENTS = {"Exp70"}

FIGURE7B_TIL_TRIM28_NORM_CORRECTIONS = {
    "exp61_": 10.0,
}

TABLE_COLUMNS = {
    "7_1": [
        "experiment",
        "sample",
        "NT_TCF7_percent",
        "Trim28_TCF7_percent",
    ],
    "7_2": [
        "sample",
        "tissue",
        "NT_percent",
        "Trim28_percent",
        "input_NT_percent",
        "input_Trim28_percent",
        "Trim28_output_input",
    ],
    "7_3": [
        "sample",
        "treatment",
        "measurement_method",
        "dLN_NT_count",
        "dLN_Trim28_1_count",
        "dLN_Trim28_count",
        "dLN_total_count",
        "dLN_Trim28_percent",
        "TIL_NT_count",
        "TIL_Trim28_1_count",
        "TIL_Trim28_count",
        "TIL_total_count",
        "TIL_Trim28_percent",
        "input_NT_count",
        "input_Trim28_1_count",
        "input_Trim28_count",
        "input_total_count",
        "input_Trim28_percent",
        "normalized_Trim28_TIL",
        "TIL_over_dLN_Trim28_ratio",
    ],
    "7_4": [
        "record_type",
        "treatment_context",
        "mouse_id",
        "group",
        "day",
        "tumor_volume_mm3",
        "event",
        "event_label",
        "end_reason",
    ],
    "7_5": [
        "sample",
        "mouse_key",
        "site",
        "NT_count",
        "Trim28_1_count",
        "Trim28_count",
        "total_count",
        "NT_fraction",
        "Trim28_fraction",
        "input_NT_count",
        "input_Trim28_1_count",
        "input_Trim28_count",
        "input_total_count",
        "input_Trim28_fraction",
        "Trim28_output_input",
    ],
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--table", choices=sorted(FIGURE7_FILES), required=True)
    parser.add_argument("--clean-dir", required=True, type=Path)
    parser.add_argument("--parsed-dir", type=Path)
    parser.add_argument("--output", required=True, type=Path)
    return parser.parse_args()


def read_clean_table(clean_dir: Path, filename: str, required_columns: list[str]) -> pd.DataFrame:
    path = clean_dir / filename
    df = pd.read_csv(path, dtype=str, keep_default_na=False)
    missing = sorted(set(required_columns) - set(df.columns))
    if missing:
        missing_str = ", ".join(missing)
        raise ValueError(f"Missing columns in {path}: {missing_str}")
    return df


def write_table(df: pd.DataFrame, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(output_path, sep="\t", index=False)
    print(f"Saved {output_path} with {len(df)} rows")


def normalize_treatment_label(value: object) -> object:
    if not isinstance(value, str):
        return value
    if value == "PD1":
        return "anti-PD-1"
    if value == "ISO":
        return "isotype control"
    return value.replace("anti-PD1", "anti-PD-1")


def format_ratio(value: float) -> str:
    return f"{value:.12g}"


def sample_label(experiment: str, sample: str) -> str:
    cleaned_sample = str(sample).strip()
    if cleaned_sample.startswith(experiment):
        return cleaned_sample
    return f"{experiment}_{cleaned_sample}"


def percent_from_fraction(value: str) -> str:
    if value == "":
        return ""
    return format_ratio(float(value) * 100.0)


def is_blank(value: object) -> bool:
    return value is None or str(value) == ""


def to_float(value: object) -> float | None:
    if is_blank(value):
        return None
    return float(value)


def format_optional_number(value: float | int | str | None) -> str:
    if is_blank(value):
        return ""
    return format_ratio(float(value))


def format_count(value: float | int | str | None) -> str:
    if is_blank(value):
        return ""
    number = float(value)
    if number.is_integer():
        return str(int(number))
    return format_ratio(number)


def apply_figure7b_til_trim28_norm_corrections(
    table: pd.DataFrame,
    column: str = "TIL_Trim28_norm",
) -> pd.DataFrame:
    """Correct source-workbook decimal-place error in Exp61 tumor Trim28 normalization."""
    table = table.copy()
    if "sample" not in table.columns or column not in table.columns:
        return table

    values = pd.to_numeric(table[column], errors="coerce")
    for sample_prefix, multiplier in FIGURE7B_TIL_TRIM28_NORM_CORRECTIONS.items():
        mask = table["sample"].astype(str).str.startswith(sample_prefix) & values.notna()
        table.loc[mask, column] = values.loc[mask] * multiplier
    return table


def count_fields(record: dict[str, object] | None, prefix: str) -> dict[str, str]:
    if record is None:
        return {
            f"{prefix}_NT_count": "",
            f"{prefix}_Trim28_1_count": "",
            f"{prefix}_Trim28_count": "",
            f"{prefix}_total_count": "",
        }
    return {
        f"{prefix}_NT_count": format_count(record.get("NT_count")),
        f"{prefix}_Trim28_1_count": format_count(record.get("Trim28_1_count")),
        f"{prefix}_Trim28_count": format_count(record.get("Trim28_count")),
        f"{prefix}_total_count": format_count(record.get("total_count")),
    }


def exp74_guide_count_lookup(parsed_dir: Path) -> dict[str, dict[str, object]]:
    workbook = parsed_dir.parent / "z_drive" / "guide_counts2.xlsx"
    wb = load_workbook(workbook, read_only=True, data_only=True)
    ws = wb["counts2"]
    counts: dict[str, dict[str, object]] = {}
    for row in ws.iter_rows(min_row=5, values_only=True):
        sample = row[0]
        if sample is None:
            continue
        nt_count = row[1]
        trim28_1_count = row[2]
        trim28_2_count = row[3]
        if nt_count is None and trim28_1_count is None and trim28_2_count is None:
            continue
        try:
            nt = float(nt_count or 0)
            trim28_1 = float(trim28_1_count or 0)
            trim28_2 = float(trim28_2_count or 0)
        except (TypeError, ValueError):
            continue
        trim28 = trim28_1 + trim28_2
        total = nt + trim28
        counts[str(sample)] = {
            "NT_count": nt_count,
            "Trim28_1_count": trim28_1_count,
            "Trim28_2_count": trim28_2_count,
            "Trim28_count": trim28,
            "total_count": total,
            "NT_fraction": nt / total if total else "",
            "Trim28_fraction": trim28 / total if total else "",
        }
    return counts


def exp74_count_sample(mouse: str, site: str) -> str:
    mouse_text = str(mouse).strip()
    mouse_id = mouse_text.split("mouse", 1)[1].strip() if "mouse" in mouse_text else mouse_text
    if site == "dLN":
        return f"Exp 74_{mouse_id} dLN"
    if site == "TIL":
        return f"Exp 74_{mouse_id} TIL"
    if site == "Peritumor skin":
        return f"Exp 74_{mouse_id}A"
    if site == "Distant skin":
        return f"Exp 74_{mouse_id}C"
    raise ValueError(f"Unsupported Exp74 count site: {site}")


def exp66_exp70_guide_count_lookup(parsed_dir: Path) -> dict[str, dict[str, object]]:
    source = read_clean_table(
        parsed_dir,
        "fig7e_exp70_exp66_skin_til_guide_counts.csv",
        ["sample", "NT_count", "Trim28_count", "NT_fraction", "Trim28_fraction"],
    )
    lookup: dict[str, dict[str, object]] = {}
    for _, row in source.iterrows():
        nt_count = to_float(row["NT_count"])
        trim28_count = to_float(row["Trim28_count"])
        total_count = (
            nt_count + trim28_count
            if nt_count is not None and trim28_count is not None
            else None
        )
        lookup[str(row["sample"])] = {
            "NT_count": row["NT_count"],
            "Trim28_1_count": "",
            "Trim28_2_count": "",
            "Trim28_count": row["Trim28_count"],
            "total_count": total_count,
            "NT_fraction": row["NT_fraction"],
            "Trim28_fraction": row["Trim28_fraction"],
        }
    return lookup


def simple_source_table(clean_dir: Path, table_id: str) -> pd.DataFrame:
    filename = FIGURE7_FILES[table_id]
    if not isinstance(filename, str):
        raise TypeError(f"Expected a single filename for {table_id}")
    df = read_clean_table(clean_dir, filename, TABLE_COLUMNS[table_id])
    return df.loc[:, TABLE_COLUMNS[table_id]].copy()


def table_7_2(clean_dir: Path, parsed_dir: Path | None) -> pd.DataFrame:
    if parsed_dir is None:
        raise ValueError("--parsed-dir is required when building table 7_2")

    clean = read_clean_table(
        clean_dir,
        "Fig7B_competitive_recovery_excluding_pd1_overlap_clean.csv",
        ["sample", "experiment", "dLN_Trim28_norm", "TIL_Trim28_norm"],
    )
    clean = apply_figure7b_til_trim28_norm_corrections(clean)
    source = read_clean_table(
        parsed_dir,
        "fig7b_current_compilation_ratios.csv",
        [
            "sample",
            "experiment",
            "dLN_NT_percent",
            "dLN_Trim28_percent",
            "TIL_NT_percent",
            "TIL_Trim28_percent",
            "dLN_NT_norm",
            "dLN_Trim28_norm",
            "TIL_NT_norm",
            "TIL_Trim28_norm",
        ],
    )
    source = apply_figure7b_til_trim28_norm_corrections(source)
    source = source.loc[
        :,
        [
            "sample",
            "experiment",
            "dLN_NT_percent",
            "dLN_Trim28_percent",
            "TIL_NT_percent",
            "TIL_Trim28_percent",
            "dLN_NT_norm",
            "dLN_Trim28_norm",
            "TIL_NT_norm",
            "TIL_Trim28_norm",
        ],
    ].copy()
    source = source.rename(
        columns={
            "dLN_NT_norm": "source_dLN_NT_norm",
            "dLN_Trim28_norm": "source_dLN_Trim28_norm",
            "TIL_NT_norm": "source_TIL_NT_norm",
            "TIL_Trim28_norm": "source_TIL_Trim28_norm",
        }
    )

    df = clean.merge(
        source,
        on=["sample", "experiment"],
        how="left",
        validate="one_to_one",
    )
    missing_source = (
        df["dLN_NT_percent"].eq("")
        & df["dLN_Trim28_percent"].eq("")
        & df["TIL_NT_percent"].eq("")
        & df["TIL_Trim28_percent"].eq("")
    )
    if missing_source.any():
        missing_samples = ", ".join(df.loc[missing_source, "sample"].tolist())
        raise ValueError(f"Missing raw percent source rows for Table 7_2 samples: {missing_samples}")

    fallback_inputs: dict[tuple[str, str, str], float] = {}
    for experiment, exp_df in df.groupby("experiment"):
        for tissue in ["dLN", "TIL"]:
            for group in ["NT", "Trim28"]:
                values: list[float] = []
                percent_col = f"{tissue}_{group}_percent"
                norm_col = f"source_{tissue}_{group}_norm"
                for _, row in exp_df.iterrows():
                    percent = to_float(row[percent_col])
                    norm = to_float(row[norm_col])
                    if percent is not None and norm not in (None, 0):
                        values.append(percent / norm)
                if values:
                    fallback_inputs[(experiment, tissue, group)] = sorted(values)[len(values) // 2]

    rows: list[dict[str, str]] = []
    for _, row in df.iterrows():
        for source_tissue, plotted_tissue in [("dLN", "dLN"), ("TIL", "tumor")]:
            plotted_value = to_float(row[f"{source_tissue}_Trim28_norm"])
            if plotted_value is None:
                continue

            output_values: dict[str, str] = {}
            input_values: dict[str, str] = {}
            for group in ["NT", "Trim28"]:
                percent_col = f"{source_tissue}_{group}_percent"
                norm_col = f"source_{source_tissue}_{group}_norm"
                percent = to_float(row[percent_col])
                norm = to_float(row[norm_col])
                input_percent = (
                    percent / norm
                    if percent is not None and norm not in (None, 0)
                    else fallback_inputs.get((row["experiment"], source_tissue, group))
                )
                if percent is None and input_percent is not None and norm is not None:
                    percent = norm * input_percent
                output_values[group] = format_optional_number(percent)
                input_values[group] = format_optional_number(input_percent)

            rows.append(
                {
                    "sample": str(row["sample"]),
                    "tissue": plotted_tissue,
                    "NT_percent": output_values["NT"],
                    "Trim28_percent": output_values["Trim28"],
                    "input_NT_percent": input_values["NT"],
                    "input_Trim28_percent": input_values["Trim28"],
                    "Trim28_output_input": format_optional_number(plotted_value),
                }
            )

    return pd.DataFrame(rows, columns=TABLE_COLUMNS["7_2"])


def table_7_3(clean_dir: Path, parsed_dir: Path | None) -> pd.DataFrame:
    if parsed_dir is None:
        raise ValueError("--parsed-dir is required when building table 7_3")

    clean = read_clean_table(
        clean_dir,
        "Fig7C_PD1_enrichment_plus_exp74_mannwhitney_clean.csv",
        [
            "experiment",
            "sample",
            "source_type",
            "treatment_for_plot",
            "normalized_Trim28_TIL",
            "TIL_over_dLN_Trim28_ratio",
        ],
    )
    flow = read_clean_table(
        clean_dir,
        "Fig7C_PD1_enrichment_curated_flow_clean.csv",
        ["experiment", "sample", "dLN_Trim28_percent", "TIL_Trim28_percent"],
    )
    exp74 = read_clean_table(
        parsed_dir,
        "fig7bce_exp74_guide_counts_summary.csv",
        ["mouse", "dLN_Trim28_fraction", "TIL_Trim28_fraction"],
    )
    exp74_counts = exp74_guide_count_lookup(parsed_dir)

    flow_lookup = {
        (str(row["experiment"]), str(row["sample"])): row
        for _, row in flow.iterrows()
    }
    exp74_lookup = {
        f"mouse {str(row['mouse']).strip()}": row
        for _, row in exp74.iterrows()
    }

    rows: list[dict[str, str]] = []
    for _, row in clean.iterrows():
        experiment = str(row["experiment"])
        sample = str(row["sample"])
        source_type = str(row["source_type"])
        normalized = str(row["normalized_Trim28_TIL"])

        if source_type == "flow cytometry":
            raw = flow_lookup.get((experiment, sample))
            if raw is None:
                raise ValueError(f"Missing raw flow percentages for Table 7_3 row: {experiment} {sample}")
            dln_percent = str(raw["dLN_Trim28_percent"])
            til_percent = str(raw["TIL_Trim28_percent"])
            input_percent = format_ratio(float(til_percent) / float(normalized)) if normalized != "" else ""
            dln_counts = None
            til_counts = None
            input_counts = None
            measurement_method = "flow cytometry"
        elif source_type == "guide counts" and experiment == "Exp74":
            raw = exp74_lookup.get(sample)
            if raw is None:
                raise ValueError(f"Missing raw Exp74 guide-count fractions for Table 7_3 row: {sample}")
            dln_counts = exp74_counts.get(exp74_count_sample(sample, "dLN"))
            til_counts = exp74_counts.get(exp74_count_sample(sample, "TIL"))
            input_counts = exp74_counts.get("Exp74input")
            if dln_counts is None or til_counts is None or input_counts is None:
                raise ValueError(f"Missing Exp74 guide counts for Table 7_3 row: {sample}")
            dln_percent = percent_from_fraction(str(raw["dLN_Trim28_fraction"]))
            til_percent = percent_from_fraction(str(raw["TIL_Trim28_fraction"]))
            input_percent = percent_from_fraction(format_ratio(input_counts["Trim28_fraction"]))
            measurement_method = "gRNA sequencing"
        else:
            raise ValueError(f"Unsupported Table 7_3 source row: {experiment} {sample} {source_type}")

        table_row = {
            "sample": sample_label(experiment, sample),
            "treatment": normalize_treatment_label(str(row["treatment_for_plot"])),
            "measurement_method": measurement_method,
            "dLN_Trim28_percent": dln_percent,
            "TIL_Trim28_percent": til_percent,
            "input_Trim28_percent": input_percent,
            "normalized_Trim28_TIL": normalized,
            "TIL_over_dLN_Trim28_ratio": str(row["TIL_over_dLN_Trim28_ratio"]),
        }
        table_row.update(count_fields(dln_counts, "dLN"))
        table_row.update(count_fields(til_counts, "TIL"))
        table_row.update(count_fields(input_counts, "input"))
        rows.append(table_row)

    return pd.DataFrame(rows, columns=TABLE_COLUMNS["7_3"])


def add_growth_rows(rows: list[dict[str, str]], df: pd.DataFrame, treatment_context: str) -> None:
    for _, row in df.iterrows():
        rows.append(
            {
                "record_type": "tumor_volume",
                "treatment_context": treatment_context,
                "mouse_id": row["mouse_id"],
                "group": row["group"],
                "day": row["day"],
                "tumor_volume_mm3": row["tumor_volume_mm3"],
                "event": "",
                "event_label": "",
                "end_reason": "",
            }
        )


def add_survival_rows(rows: list[dict[str, str]], df: pd.DataFrame, treatment_context: str) -> None:
    for _, row in df.iterrows():
        rows.append(
            {
                "record_type": "survival_endpoint",
                "treatment_context": treatment_context,
                "mouse_id": row["mouse_id"],
                "group": row["group"],
                "day": row["time_day"],
                "tumor_volume_mm3": "",
                "event": row["event"],
                "event_label": row["event_label"],
                "end_reason": row["end_reason"],
            }
        )


def table_7_4(clean_dir: Path) -> pd.DataFrame:
    exp71_growth = read_clean_table(
        clean_dir,
        "Fig7D_exp71_growth_no_pd1_forwardfilled_clean.csv",
        ["experiment", "mouse_id", "group", "day", "tumor_volume_mm3", "note"],
    )
    exp71_survival = read_clean_table(
        clean_dir,
        "Fig7D_exp71_survival_no_pd1_forwardfilled_clean.csv",
        ["experiment", "mouse_id", "group", "time_day", "event", "event_label", "end_reason", "note"],
    )
    exp67_growth = read_clean_table(
        clean_dir,
        "Fig7D_exp67_growth_pd1_clean.csv",
        ["experiment", "mouse_id", "group", "day", "tumor_volume_mm3", "note"],
    )
    exp67_survival = read_clean_table(
        clean_dir,
        "Fig7D_exp67_survival_pd1_clean.csv",
        ["experiment", "mouse_id", "group", "time_day", "event", "event_label", "end_reason", "note"],
    )
    rows: list[dict[str, str]] = []
    add_growth_rows(rows, exp71_growth, "isotype/no anti-PD-1 documented")
    add_survival_rows(rows, exp71_survival, "isotype/no anti-PD-1 documented")
    add_growth_rows(rows, exp67_growth, "anti-PD-1")
    add_survival_rows(rows, exp67_survival, "anti-PD-1")

    return pd.DataFrame(rows, columns=TABLE_COLUMNS["7_4"])


def skin_output_input(experiment: str, trim28_fraction: str) -> str:
    if experiment not in SKIN_INPUT_TRIM28_FRACTIONS:
        raise ValueError(f"No input Trim28 fraction for {experiment}")
    if trim28_fraction == "":
        return ""
    return format_ratio(float(trim28_fraction) / SKIN_INPUT_TRIM28_FRACTIONS[experiment])


def table_7_5(clean_dir: Path, parsed_dir: Path | None) -> pd.DataFrame:
    if parsed_dir is None:
        raise ValueError("--parsed-dir is required when building table 7_5")

    included = read_clean_table(
        clean_dir,
        "Fig7E_skin_distribution_3bin_clean.csv",
        ["experiment", "sample", "mouse_key", "Trim28_fraction", "site"],
    )
    pairs = read_clean_table(
        clean_dir,
        "Fig7E_peritumor_vs_distant_pairs.csv",
        ["experiment", "mouse_key"],
    )

    pairs = pairs.loc[~pairs["experiment"].astype(str).isin(SKIN_EXCLUDED_EXPERIMENTS)].copy()
    included = included.loc[
        ~included["experiment"].astype(str).isin(SKIN_EXCLUDED_EXPERIMENTS)
    ].copy()
    paired_mouse_keys = set(pairs["mouse_key"])
    exp66_exp70_counts = exp66_exp70_guide_count_lookup(parsed_dir)
    exp74_counts = exp74_guide_count_lookup(parsed_dir)
    exp74_input = exp74_counts.get("Exp74input")
    if exp74_input is None:
        raise ValueError("Missing Exp74 input counts")

    rows: list[dict[str, str]] = []
    plotted = included.loc[included["mouse_key"].isin(paired_mouse_keys)].copy()
    for _, row in plotted.iterrows():
        experiment = row["experiment"]
        if experiment == "Exp74":
            counts = exp74_counts.get(exp74_count_sample(row["sample"], row["site"]))
            input_counts = exp74_input
        else:
            counts = exp66_exp70_counts.get(str(row["sample"]))
            input_counts = None
        if counts is None:
            raise ValueError(f"Missing guide counts for Table 7_5 row: {row['sample']} {row['site']}")

        input_trim28_fraction = SKIN_INPUT_TRIM28_FRACTIONS[experiment]
        output_nt_fraction = counts.get("NT_fraction")
        output_trim28_fraction = row["Trim28_fraction"]

        table_row = {
            "sample": row["sample"],
            "mouse_key": row["mouse_key"],
            "site": row["site"],
            "NT_count": format_count(counts.get("NT_count")),
            "Trim28_1_count": format_count(counts.get("Trim28_1_count")),
            "Trim28_count": format_count(counts.get("Trim28_count")),
            "total_count": format_count(counts.get("total_count")),
            "NT_fraction": format_optional_number(output_nt_fraction),
            "Trim28_fraction": output_trim28_fraction,
            "input_Trim28_fraction": format_ratio(input_trim28_fraction),
            "Trim28_output_input": skin_output_input(experiment, output_trim28_fraction),
        }
        table_row.update(count_fields(input_counts, "input"))
        rows.append(table_row)

    return pd.DataFrame(rows, columns=TABLE_COLUMNS["7_5"])


def build_table(table_id: str, clean_dir: Path, parsed_dir: Path | None) -> pd.DataFrame:
    if table_id == "7_1":
        return simple_source_table(clean_dir, table_id)
    if table_id == "7_3":
        return table_7_3(clean_dir, parsed_dir)
    if table_id == "7_2":
        return table_7_2(clean_dir, parsed_dir)
    if table_id == "7_4":
        return table_7_4(clean_dir)
    if table_id == "7_5":
        return table_7_5(clean_dir, parsed_dir)
    raise ValueError(f"Unsupported table id: {table_id}")


def main() -> None:
    args = parse_args()
    table = build_table(args.table, args.clean_dir, args.parsed_dir)
    write_table(table, args.output)


if __name__ == "__main__":
    main()
