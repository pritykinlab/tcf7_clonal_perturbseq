#!/usr/bin/env python3
"""Build a supplementary table workbook from TSV files and caption text."""

from __future__ import annotations

import argparse
import csv
import re
import sys
from pathlib import Path
from typing import NamedTuple

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


CAPTION_RE = re.compile(
    r"^Table\s+((?:S\d+(?:\.\d+[A-Za-z]?)*)|(?:M\d+[A-Za-z]?))\.\s*(.+)$"
)
INT_RE = re.compile(r"^[+-]?(0|[1-9]\d*)$")
FLOAT_RE = re.compile(
    r"^[+-]?(?:(?:\d+\.\d*)|(?:\.\d+)|(?:\d+))(?:[eE][+-]?\d+)?$"
)
INVALID_SHEET_CHARS_RE = re.compile(r"[:\\/?*\[\]]")
MAX_SHEET_NAME_LEN = 31


class CaptionEntry(NamedTuple):
    label: str
    stem: str
    description: str


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Create an Excel workbook with a captions sheet followed by one sheet "
            "per supplementary table TSV."
        )
    )
    parser.add_argument(
        "--captions",
        required=True,
        type=Path,
        help="Caption text file with paragraphs beginning 'Table S...'.",
    )
    parser.add_argument(
        "--output",
        required=True,
        type=Path,
        help="Output .xlsx workbook path.",
    )
    parser.add_argument(
        "tables",
        nargs="+",
        type=Path,
        help="Supplementary table TSV files to include.",
    )
    return parser.parse_args()


def label_to_stem(label: str) -> str:
    if label.startswith("S"):
        return label[1:].replace(".", "_")
    if label.startswith("M"):
        return label
    raise ValueError(f"Expected caption label to start with 'S' or 'M': {label}")


def read_caption_entries(caption_path: Path) -> list[CaptionEntry]:
    paragraphs = [
        re.sub(r"\s+", " ", paragraph.strip())
        for paragraph in re.split(r"\n\s*\n", caption_path.read_text().strip())
        if paragraph.strip()
    ]
    entries: list[CaptionEntry] = []
    seen_labels: set[str] = set()
    for paragraph in paragraphs:
        match = CAPTION_RE.match(paragraph)
        if match is None:
            raise ValueError(
                "Each caption paragraph must begin with a label like "
                f"'Table S1.1.' or 'Table M1.': {paragraph[:120]}"
            )
        label, description = match.groups()
        if label in seen_labels:
            raise ValueError(f"Duplicate caption label: {label}")
        seen_labels.add(label)
        entries.append(
            CaptionEntry(
                label=label,
                stem=label_to_stem(label),
                description=description,
            )
        )
    return entries


def order_table_paths(
    tables: list[Path], captions: list[CaptionEntry]
) -> list[tuple[CaptionEntry, Path]]:
    table_by_stem: dict[str, Path] = {}
    for table in tables:
        if table.suffix.lower() != ".tsv":
            raise ValueError(f"Supplementary table input is not a TSV: {table}")
        if table.stem in table_by_stem:
            raise ValueError(f"Duplicate supplementary table stem: {table.stem}")
        table_by_stem[table.stem] = table

    caption_by_stem = {entry.stem: entry for entry in captions}
    missing_captions = sorted(set(table_by_stem) - set(caption_by_stem))
    missing_tables = sorted(set(caption_by_stem) - set(table_by_stem))
    messages = []
    if missing_captions:
        messages.append(
            "TSV files are missing captions: "
            + ", ".join(f"{stem}.tsv" for stem in missing_captions)
        )
    if missing_tables:
        messages.append(
            "Captions are missing TSV files: "
            + ", ".join(f"{caption_by_stem[stem].label}" for stem in missing_tables)
        )
    if messages:
        raise ValueError("; ".join(messages))

    return [(entry, table_by_stem[entry.stem]) for entry in captions]


def coerce_cell(value: str):
    candidate = value.strip()
    if candidate == "":
        return None
    if candidate != value:
        return value
    if INT_RE.match(candidate):
        return int(candidate)
    if FLOAT_RE.match(candidate) and ("." in candidate or "e" in candidate.lower()):
        return float(candidate)
    return value


def make_cell(
    worksheet,
    value,
    *,
    bold: bool = False,
    wrap: bool = False,
    vertical: str = "top",
) -> WriteOnlyCell:
    cell = WriteOnlyCell(worksheet, value=value)
    if bold:
        cell.font = Font(bold=True)
    if wrap or vertical != "bottom":
        cell.alignment = Alignment(wrap_text=wrap, vertical=vertical)
    return cell


def safe_sheet_title(base_title: str, used_titles: set[str]) -> str:
    title = INVALID_SHEET_CHARS_RE.sub("_", base_title).strip() or "sheet"
    title = title[:MAX_SHEET_NAME_LEN]
    if title not in used_titles:
        used_titles.add(title)
        return title

    suffix_index = 2
    while True:
        suffix = f"_{suffix_index}"
        candidate = title[: MAX_SHEET_NAME_LEN - len(suffix)] + suffix
        if candidate not in used_titles:
            used_titles.add(candidate)
            return candidate
        suffix_index += 1


def write_captions_sheet(workbook: Workbook, entries: list[CaptionEntry]) -> None:
    worksheet = workbook.create_sheet("captions")
    worksheet.column_dimensions["A"].width = 14
    worksheet.column_dimensions["B"].width = 120
    worksheet.append(
        [
            make_cell(worksheet, "Table", bold=True),
            make_cell(worksheet, "Caption", bold=True),
        ]
    )
    for entry in entries:
        worksheet.append(
            [
                make_cell(worksheet, entry.label, bold=True),
                make_cell(worksheet, entry.description, wrap=True),
            ]
        )


def write_table_sheet(
    workbook: Workbook,
    sheet_title: str,
    table_path: Path,
    used_titles: set[str],
) -> None:
    worksheet = workbook.create_sheet(safe_sheet_title(sheet_title, used_titles))
    with table_path.open(newline="") as handle:
        reader = csv.reader(handle, delimiter="\t")
        try:
            header = next(reader)
        except StopIteration as exc:
            raise ValueError(f"Supplementary table is empty: {table_path}") from exc

        for column_number, column_name in enumerate(header, start=1):
            width = min(max(len(column_name) + 2, 12), 35)
            column_letter = get_column_letter(column_number)
            worksheet.column_dimensions[column_letter].width = width
        worksheet.append([make_cell(worksheet, value, bold=True) for value in header])

        for row in reader:
            worksheet.append([coerce_cell(value) for value in row])


def build_workbook(
    caption_path: Path,
    table_paths: list[Path],
    output_path: Path,
) -> None:
    captions = read_caption_entries(caption_path)
    ordered_tables = order_table_paths(table_paths, captions)

    workbook = Workbook(write_only=True)
    write_captions_sheet(workbook, captions)

    used_titles = {"captions"}
    for entry, table_path in ordered_tables:
        write_table_sheet(workbook, entry.label, table_path, used_titles)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def main() -> int:
    args = parse_args()
    try:
        build_workbook(args.captions, args.tables, args.output)
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
